"""Figura 7 en Excel: abanico del patrimonio (P10-P90, P25-P75 y mediana), un panel por escenario.
Usa los percentiles de estado.pkl (analisis_complementario.py). Uso: python3 figura_abanico_excel.py"""
import os
import openpyxl
import pandas as pd
from openpyxl.chart import AreaChart, LineChart, Reference
from openpyxl.chart.legend import LegendEntry
from openpyxl.styles import Font

BASE = os.path.join(os.path.dirname(__file__), "..")
T = pd.read_pickle(os.path.join(BASE, "resultados/estado.pkl"))["T"]
AZUL, CLARO, MEDIO = "2A78D6", "D6E4F7", "A9C6EE"

wb = openpyxl.Workbook(); wb.remove(wb.active)
for hoja, k, tit in [("Histórico", "Abanico_historico", "Retornos históricos 2013-2026"),
                     ("Prospectivo", "Abanico_prospectivo", "Supuestos prospectivos")]:
    a = T[k]; ws = wb.create_sheet(hoja)
    hdr = ["Año", "P10", "P25", "P50 (mediana)", "P75", "P90", "", "Base (P10)", "P10-P25", "P25-P75", "P75-P90"]
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True)
    for i, r in enumerate(a.itertuples(), start=2):
        ws.append([r.Año] + [round(v / 1e3, 1) for v in (r.P10, r.P25, r.P50, r.P75, r.P90)] + [None,
                  f"=B{i}", f"=C{i}-B{i}", f"=E{i}-C{i}", f"=F{i}-E{i}"])
    n = len(a) + 1
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=11):
        for c in row: c.number_format = "#,##0"
    ws.cell(row=n + 2, column=1, value="Miles de USD de hoy. Percentiles del patrimonio al cierre de cada año en 20.000 escenarios del Montecarlo oficial. Las columnas H-K son las franjas que se apilan en el gráfico (la base va sin relleno). Fuente: codigo/analisis_complementario.py.")

    area = AreaChart(); area.grouping = "stacked"
    area.add_data(Reference(ws, min_col=8, max_col=11, min_row=1, max_row=n), titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n); area.set_categories(cats)
    for s, fill in zip(area.series, [None, CLARO, MEDIO, CLARO]):
        if fill is None: s.graphicalProperties.noFill = True
        else: s.graphicalProperties.solidFill = fill
        s.graphicalProperties.line.noFill = True
    line = LineChart()
    line.add_data(Reference(ws, min_col=4, min_row=1, max_row=n), titles_from_data=True)
    line.set_categories(cats)
    line.series[0].graphicalProperties.line.solidFill = AZUL; line.series[0].graphicalProperties.line.width = 28575
    line.series[0].smooth = False
    area += line
    area.title = f"Abanico del patrimonio — {tit}"
    area.y_axis.title = "Patrimonio (miles de USD de hoy)"; area.x_axis.title = "Año"
    area.y_axis.number_format = "#,##0"; area.y_axis.scaling.min = 0
    area.y_axis.delete = False; area.x_axis.delete = False
    area.x_axis.tickLblSkip = 5; area.x_axis.tickMarkSkip = 5
    area.legend.position = "b"; area.legend.legendEntry = [LegendEntry(idx=0, delete=True)]  # oculta la base
    area.width = 22; area.height = 12
    ws.add_chart(area, "M2")
    ws.column_dimensions["A"].width = 8
wb.save(os.path.join(BASE, "informe/figuras/Abanico_patrimonio.xlsx"))
