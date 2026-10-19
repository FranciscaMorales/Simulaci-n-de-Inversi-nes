"""Todas las figuras del informe en un solo Excel (una hoja por figura, con sus datos y su gráfico nativo).
Usa estado.pkl y resultados/version_b_tir.csv. Uso: python3 figuras_excel.py
Salidas: informe/figuras/Figuras_Informe_Nexum.xlsx y informe/figuras/Robustez.xlsx (Figura 8 sola)."""
import os
import openpyxl
import pandas as pd
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font, Alignment
from config import RESULTADOS, BASE, NOMBRES, ACTIVOS

S = pd.read_pickle(os.path.join(RESULTADOS, "estado.pkl")); T = S["T"]
OUT = os.path.join(BASE, "informe", "figuras")
AZUL, NARANJA, AQUA, AMARILLO, CLARO, MEDIO = "2A78D6", "EB6834", "1BAF7A", "EDA100", "D6E4F7", "A9C6EE"
COL = [AZUL, NARANJA, AQUA, AMARILLO]
B = Font(bold=True)

def encabezado(ws, titulo, nota, cols):
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = nota; ws["A2"].alignment = Alignment(wrap_text=False)
    for j, c in enumerate(cols, 1): ws.cell(row=4, column=j, value=c).font = B
    return 5

def ejes(ch, xt=None, yt=None, yfmt=None):
    ch.x_axis.delete = False; ch.y_axis.delete = False
    if xt: ch.x_axis.title = xt
    if yt: ch.y_axis.title = yt
    if yfmt: ch.y_axis.number_format = yfmt
    ch.width, ch.height = 24, 12
    return ch

def barras_h(ws, titulo, cats, cols_val, r0, n, colores, fmt="0.0%", etiquetas=True):
    ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.gapWidth = 60; ch.title = titulo
    for k, c in enumerate(cols_val):
        ch.add_data(Reference(ws, min_col=c, min_row=r0 - 1, max_row=r0 + n - 1), titles_from_data=True)
        ch.series[k].graphicalProperties.solidFill = colores[k]; ch.series[k].graphicalProperties.line.noFill = True
    ch.set_categories(Reference(ws, min_col=cats, min_row=r0, max_row=r0 + n - 1))
    ch.x_axis.scaling.orientation = "maxMin"  # primer elemento arriba, como en el informe
    if etiquetas: ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.numFmt = fmt
    return ch

def fig1(ws, version_b=False):
    t = T["TIR"].copy()
    if version_b:
        tb = round(float(pd.read_csv(os.path.join(RESULTADOS, "version_b_tir.csv"))["TIR B"][0]), 4)
        t = pd.concat([t.iloc[:1], pd.DataFrame([{"Escenario": "Versión B: casa de Villarrica en la herencia", "TIR real requerida": tb}]), t.iloc[1:]], ignore_index=True)
    r = encabezado(ws, "Figura 1%s. TIR real requerida por escenario" % (" (versión B)" if version_b else ""),
                   "Fuente: codigo/objetivos.py (TIR de los flujos de cada escenario); modelo, hoja 02_Objetivos_TIR (fórmula IRR)." + (" Barra naranja: versión B (codigo/version_b.py)." if version_b else ""),
                   ["Escenario", "TIR real requerida"])
    for i, x in enumerate(t.itertuples()): ws.cell(row=r + i, column=1, value=x.Escenario); ws.cell(row=r + i, column=2, value=round(float(x._2), 4)).number_format = "0.00%"
    ch = barras_h(ws, "Retorno real requerido para cumplir los tres objetivos", 1, [2], r, len(t), [AZUL], "0.00%")
    if version_b: p = DataPoint(idx=1); p.graphicalProperties.solidFill = NARANJA; ch.series[0].dPt.append(p)
    ch.legend = None; ws.add_chart(ejes(ch, yt="TIR real anual requerida", yfmt="0%"), "D4"); ws.column_dimensions["A"].width = 48

def fig2(ws):
    a = T["Asignacion"]; cols = [("Pesos aprobados (en implementación)", "Black-Litterman (aprobado)"), ("Prior paridad de riesgo", "Prior de paridad de riesgo"), ("Markowitz histórico", "Markowitz histórico")]
    r = encabezado(ws, "Figura 2. Asignación estratégica: Black-Litterman frente a sus referencias",
                   "Fuente: codigo/portafolio.py; modelo, hojas 05 (Black-Litterman) y 06 (asignación).", ["Activo"] + [n for _, n in cols])
    for i, x in enumerate(ACTIVOS):
        ws.cell(row=r + i, column=1, value=NOMBRES[x])
        for j, (c, _) in enumerate(cols): ws.cell(row=r + i, column=2 + j, value=round(float(a[c].iloc[i]), 4)).number_format = "0.0%"
    ch = barras_h(ws, "Asignación estratégica: Black-Litterman frente a sus referencias", 1, [2, 3, 4], r, len(ACTIVOS), COL, etiquetas=False)
    ch.legend.position = "b"; ws.add_chart(ejes(ch, yt="Peso", yfmt="0%"), "F4"); ws.column_dimensions["A"].width = 34

def fig3(ws):
    rc = T["Contribucion_riesgo"]
    r = encabezado(ws, "Figura 3. Peso frente a contribución al riesgo total",
                   "Contribución = peso × covarianza con la cartera / varianza de la cartera (pesos aprobados, covarianza 2013-2026). Fuente: codigo/analisis_complementario.py; modelo, hoja 06.",
                   ["Activo", "Peso en la cartera", "Contribución al riesgo"])
    for i, x in enumerate(rc.itertuples()):
        ws.cell(row=r + i, column=1, value=x.Activo)
        ws.cell(row=r + i, column=2, value=round(float(rc["Peso"].iloc[i]), 4)).number_format = "0.0%"
        ws.cell(row=r + i, column=3, value=round(float(rc["Contribución al riesgo (%)"].iloc[i]), 4)).number_format = "0.0%"
    ch = barras_h(ws, "Peso frente a contribución al riesgo total de la cartera", 1, [2, 3], r, len(rc), [AZUL, NARANJA])
    ch.legend.position = "b"; ws.add_chart(ejes(ch, yfmt="0%"), "E4"); ws.column_dimensions["A"].width = 34

def serie_crecimiento():
    rr = pd.DataFrame({"Cartera Nexum": S["port"], "Benchmark de política": S["bench"], "AGF actual 85/15": S["agf"], "S&P 500": S["df"].SPY}).loc["2016-01-31":]
    return 100 * (1 + rr).cumprod()

def lineas(ws, titulo, r, n, ncols, colores, yt, yfmt):
    ch = LineChart(); ch.title = titulo
    ch.add_data(Reference(ws, min_col=2, max_col=1 + ncols, min_row=r - 1, max_row=r + n - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r, max_row=r + n - 1))
    for k, s in enumerate(ch.series):
        s.graphicalProperties.line.solidFill = colores[k]; s.graphicalProperties.line.width = 28575 if k == 0 else 19050; s.smooth = False
    ch.x_axis.number_format = "mmm-yy"; ch.x_axis.tickLblSkip = 12; ch.x_axis.tickMarkSkip = 12
    ch.legend.position = "b"; return ejes(ch, yt=yt, yfmt=yfmt)

def fig4(ws):
    w = serie_crecimiento()
    r = encabezado(ws, "Figura 4. Crecimiento de USD 100 (ene-2016 a ago-2026, USD nominal)",
                   "Fuente: codigo/portafolio.py (retornos mensuales de la cartera, el benchmark y la AGF); resultados/retornos_cartera_benchmark.csv.", ["Fecha"] + list(w.columns))
    for i, (d, row) in enumerate(w.iterrows()):
        ws.cell(row=r + i, column=1, value=d.date()).number_format = "mmm-yy"
        for j, v in enumerate(row): ws.cell(row=r + i, column=2 + j, value=round(float(v), 2)).number_format = "0.0"
    ws.add_chart(lineas(ws, "Crecimiento de USD 100", r, len(w), 4, COL, "Valor de USD 100 invertidos", "0"), "G4"); ws.column_dimensions["A"].width = 11

def fig5(ws):
    w = serie_crecimiento()[["Cartera Nexum", "AGF actual 85/15"]]; dd = w / w.cummax() - 1
    r = encabezado(ws, "Figura 5. Caídas desde el máximo: Nexum frente a la AGF",
                   "Caída = valor / máximo anterior − 1, sobre el crecimiento de USD 100 de la Figura 4.", ["Fecha"] + list(dd.columns))
    for i, (d, row) in enumerate(dd.iterrows()):
        ws.cell(row=r + i, column=1, value=d.date()).number_format = "mmm-yy"
        for j, v in enumerate(row): ws.cell(row=r + i, column=2 + j, value=round(float(v), 4)).number_format = "0.0%"
    ws.add_chart(lineas(ws, "Caídas desde el máximo", r, len(dd), 2, [AZUL, AQUA], "Caída desde el máximo", "0%"), "E4"); ws.column_dimensions["A"].width = 11

def fig6(ws):
    mc = T["Montecarlo"]; objs = ["P(Educación)", "P(Villarrica)", "P(Herencia)", "P(3 objetivos)"]
    esc = [("Nexum", "Hist", "USD", "Nexum · histórico"), ("Nexum", "Prosp", "USD", "Nexum · prospectivo"), ("AGF 85/15", "Prosp", "USD", "AGF · prospectivo"), ("Nexum", "Prosp", "UF", "Nexum · prospectivo en UF")]
    r = encabezado(ws, "Figura 6. Montecarlo: probabilidad de cumplir cada objetivo",
                   "20.000 escenarios de 50 años por bootstrap mensual. Fuente: codigo/montecarlo.py; modelo, hoja 08_Montecarlo.", ["Objetivo"] + [e[3] for e in esc])
    for i, o in enumerate(["Educación", "Villarrica", "Herencia", "Los 3 objetivos"]): ws.cell(row=r + i, column=1, value=o)
    for j, (cart, e, mon, _) in enumerate(esc):
        v = mc[(mc.Cartera == cart) & mc.Escenario.str.startswith(e) & (mc.Moneda == mon)][objs].iloc[0].values
        for i, vv in enumerate(v): ws.cell(row=r + i, column=2 + j, value=round(float(vv), 4)).number_format = "0.0%"
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.title = "Probabilidad de cumplir cada objetivo"
    ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=r - 1, max_row=r + 3), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r, max_row=r + 3))
    for k, s in enumerate(ch.series): s.graphicalProperties.solidFill = COL[k]; s.graphicalProperties.line.noFill = True
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.numFmt = "0%"
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.legend.position = "b"
    ws.add_chart(ejes(ch, yt="Probabilidad", yfmt="0%"), "G4"); ws.column_dimensions["A"].width = 16; [ws.column_dimensions[c].__setattr__("width", 22) for c in "BCDE"]

def abanico(ws, r, a, titulo, ancla):
    for i, x in enumerate(a.itertuples()):
        rr = r + i
        ws.append([x.Año] + [round(v / 1e3, 1) for v in (x.P10, x.P25, x.P50, x.P75, x.P90)] + [None, f"=B{rr}", f"=C{rr}-B{rr}", f"=E{rr}-C{rr}", f"=F{rr}-E{rr}"])
    n = len(a)
    area = AreaChart(); area.grouping = "stacked"; area.title = titulo
    area.add_data(Reference(ws, min_col=8, max_col=11, min_row=r - 1, max_row=r + n - 1), titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=r, max_row=r + n - 1); area.set_categories(cats)
    for s, fill in zip(area.series, [None, CLARO, MEDIO, CLARO]):
        if fill is None: s.graphicalProperties.noFill = True
        else: s.graphicalProperties.solidFill = fill
        s.graphicalProperties.line.noFill = True
    ln = LineChart(); ln.add_data(Reference(ws, min_col=4, min_row=r - 1, max_row=r + n - 1), titles_from_data=True); ln.set_categories(cats)
    ln.series[0].graphicalProperties.line.solidFill = AZUL; ln.series[0].graphicalProperties.line.width = 28575; ln.series[0].smooth = False
    area += ln; area.y_axis.scaling.min = 0; area.x_axis.tickLblSkip = 5; area.x_axis.tickMarkSkip = 5
    area.legend.position = "b"; area.legend.legendEntry = [LegendEntry(idx=0, delete=True)]  # oculta la base sin relleno
    ws.add_chart(ejes(area, xt="Año", yt="Patrimonio (miles de USD de hoy)", yfmt="#,##0"), ancla)

def fig7(ws):
    cols = ["Año", "P10", "P25", "P50 (mediana)", "P75", "P90", "", "Base (P10)", "P10-P25", "P25-P75", "P75-P90"]
    encabezado(ws, "Figura 7. Abanico del patrimonio real (miles de USD de hoy)",
               "Percentiles del patrimonio al cierre de cada año en los 20.000 escenarios. Columnas H-K: franjas apiladas (la base va sin relleno). Fuente: codigo/analisis_complementario.py; modelo, hoja 08_Montecarlo (B38).", [])
    ws["A4"] = "Retornos históricos 2013-2026"; ws["A4"].font = B
    ws.append(cols); [setattr(c, "font", B) for c in ws[5]]
    abanico(ws, 6, T["Abanico_historico"], "Retornos históricos 2013-2026", "M4")
    r2 = ws.max_row + 3; ws.cell(row=r2, column=1, value="Supuestos prospectivos").font = B
    for j, c in enumerate(cols, 1): ws.cell(row=r2 + 1, column=j, value=c).font = B
    ws._current_row = r2 + 1
    abanico(ws, r2 + 2, T["Abanico_prospectivo"], "Supuestos prospectivos", "M30")
    for row in ws.iter_rows(min_row=6, min_col=2, max_col=11):
        for c in row: c.number_format = "#,##0"

CORTO = {"BASE (vigente)": "Regla vigente", "BASE (vigente, 5 ETF)": "Regla vigente (5 ETF)", "Sin filtros absolutos (solo ranking)": "Sin filtros (solo ranking)", "Umbrales x2 (r6>10%, r12>20%)": "Umbrales x2", "Universo original 3 ETF (GLD, USO, DBA)": "Universo original 3 ETF", "Umbrales en cero (solo tendencia)": "Umbrales en cero", "Sin límite de volatilidad": "Sin límite de vol."}
def fig8(ws):
    rb = T["Robustez"]
    encabezado(ws, "Figura 8. Robustez: Sharpe 2013-2026 de cada variante de parámetros (naranja = regla vigente)",
               "Cada variante cambia un solo parámetro de la regla y repite el backtest completo. Fuente: codigo/validacion.py; modelo, hoja 09_Validacion.", [])
    fila = 4; anclas = ["F4", "F28", "F52"]
    for k, est in enumerate(["Momentum", "Low-Volatility", "Commodities"]):
        d = rb[rb.Estrategia == est].reset_index(drop=True)
        ws.cell(row=fila, column=1, value=est).font = B
        for j, c in enumerate(["Variante", "Sharpe", "CAGR", "Máx. caída"], 1): ws.cell(row=fila + 1, column=j, value=c).font = B
        r = fila + 2
        for i, x in d.iterrows():
            ws.cell(row=r + i, column=1, value=CORTO.get(x.Variante, x.Variante))
            ws.cell(row=r + i, column=2, value=round(float(x.Sharpe), 3)).number_format = "0.00"
            ws.cell(row=r + i, column=3, value=round(float(x.CAGR), 4)).number_format = "0.0%"
            ws.cell(row=r + i, column=4, value=round(float(x["Máx. caída"]), 4)).number_format = "0.0%"
        ch = barras_h(ws, f"{est}: Sharpe 2013-2026", 1, [2], r, len(d), [AZUL], "0.00")
        for i, v in enumerate(d.Variante):
            if v.startswith("BASE"): p = DataPoint(idx=i); p.graphicalProperties.solidFill = NARANJA; ch.series[0].dPt.append(p)
        ch.legend = None; ch.height = 9; ws.add_chart(ejes(ch, yt="Sharpe", yfmt="0.0"), anclas[k]); ch.height = 9
        fila = r + len(d) + 2
    ws.column_dimensions["A"].width = 34

FIGS = [("Fig1 TIR", fig1, "Figura 1. TIR real requerida por escenario"), ("Fig1B TIR versión B", lambda ws: fig1(ws, True), "Figura 1, versión B (barra naranja: 3,63%)"),
        ("Fig2 Asignación", fig2, "Figura 2. Asignación estratégica"), ("Fig3 Riesgo", fig3, "Figura 3. Peso vs contribución al riesgo"),
        ("Fig4 Crecimiento", fig4, "Figura 4. Crecimiento de USD 100"), ("Fig5 Caídas", fig5, "Figura 5. Caídas desde el máximo"),
        ("Fig6 Montecarlo", fig6, "Figura 6. Probabilidad por objetivo"), ("Fig7 Abanico", fig7, "Figura 7. Abanico del patrimonio"),
        ("Fig8 Robustez", fig8, "Figura 8. Robustez de parámetros")]

if __name__ == "__main__":
    wb = openpyxl.Workbook(); idx = wb.active; idx.title = "Índice"
    idx["A1"] = "Figuras del Informe Final Nexum: una hoja por figura, con los datos y el gráfico"; idx["A1"].font = Font(bold=True, size=12)
    idx["A2"] = "Generado con codigo/figuras_excel.py a partir de los mismos resultados del informe (resultados/estado.pkl)."
    for i, (hoja, f, desc) in enumerate(FIGS, start=4):
        f(wb.create_sheet(hoja)); c = idx.cell(row=i, column=1, value=desc); c.hyperlink = f"#'{hoja}'!A1"; c.style = "Hyperlink"
    idx.column_dimensions["A"].width = 60
    wb.save(os.path.join(OUT, "Figuras_Informe_Nexum.xlsx"))
    w8 = openpyxl.Workbook(); fig8(w8.active); w8.active.title = "Robustez"; w8.save(os.path.join(OUT, "Robustez.xlsx"))
    print("OK")
