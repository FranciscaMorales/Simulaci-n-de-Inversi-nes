"""Figura 1 en Excel para la versión B: TIR requerida por escenario + barra de la versión B. Uso: python3 figura_tir_b.py"""
import os
import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font

BASE = os.path.join(os.path.dirname(__file__), "..")
src = openpyxl.load_workbook(os.path.join(BASE, "informe/figuras/TIR_requerida.xlsx"))["TIR requerida"]
rows = list(src.iter_rows(min_row=2, values_only=True))
B = pd.read_csv(os.path.join(BASE, "resultados/version_b_tir.csv")).iloc[0]
rows.insert(-1, ("Versión B: casa de Villarrica en la herencia", round(float(B["TIR B"]), 4)))  # justo bajo la barra Base

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "TIR requerida"
ws.append(["Escenario", "TIR real requerida"]); ws["A1"].font = ws["B1"].font = Font(bold=True)
for r in rows: ws.append(r)
for c in ws["B"][1:]: c.number_format = "0.00%"
ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 18
n = len(rows) + 1

ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.gapWidth = 60
ch.title = "Retorno real requerido para cumplir los tres objetivos"
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=n), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
s = ch.series[0]; s.graphicalProperties.solidFill = "3F78CF"
pt = DataPoint(idx=len(rows) - 2); pt.graphicalProperties.solidFill = "E8853D"; s.dPt.append(pt)
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.numFmt = "0.00%"
ch.y_axis.title = "TIR real anual requerida"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 0.06
ch.y_axis.number_format = "0%"; ch.y_axis.delete = False; ch.x_axis.delete = False
ch.legend = None; ch.width = 26; ch.height = 12
ws.add_chart(ch, "D2")
ws.cell(row=n + 2, column=1, value="Barra naranja: versión B (la herencia se cumple con la casa de Villarrica, neta de 5% de costos de venta, más USD " + f"{B['Herencia financiera USD']:,.0f}".replace(",", ".") + " de la cartera). Fuente: codigo/version_b.py.")
wb.save(os.path.join(BASE, "informe/figuras/TIR_requerida_VersionB.xlsx"))
