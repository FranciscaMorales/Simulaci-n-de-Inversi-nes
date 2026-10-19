"""Genera la versión B del informe (la casa de Villarrica como parte de la herencia) a partir del informe oficial.
Requiere haber corrido informe.py, casa_herencia.py y version_b.py."""
import os, copy, docx, pandas as pd
from docx.shared import Pt
os.chdir(os.path.dirname(os.path.abspath(__file__)))
d = docx.Document("../informe/Informe_Final_Nexum.docx"); P = d.paragraphs
R = pd.read_csv("../resultados/version_b_montecarlo.csv"); TT = pd.read_csv("../resultados/version_b_tir.csv").iloc[0]
CH = pd.read_csv("../resultados/casa_herencia.csv")
pc = lambda v, k=1: f"{v*100:.{k}f}%".replace(".", ",")
us = lambda v: f"{v:,.0f}".replace(",", ".")
def g(c, e, m, k): return R[R.Cartera.str.startswith(c) & R.Escenario.str.startswith(e) & (R.Moneda == m)][k].iloc[0]
tb, cb, cn, hf = TT["TIR B"], TT["Cobertura B"], TT["Casa neta USD"], TT["Herencia financiera USD"]
def find(start): return next(p for p in d.paragraphs if p.text.startswith(start))
def settext(p, bold, rest):
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    if bold: p.runs[0].text = bold; p.add_run(rest)
    else: p.runs[0].text = rest
def insert_after(p, text, italic=False, style=None):
    new = copy.deepcopy(p._p); p._p.addnext(new); q = docx.text.paragraph.Paragraph(new, p._parent)
    for r in q.runs[1:]: r._r.getparent().remove(r._r)
    if not q.runs: q.add_run("")
    q.runs[0].text = text; q.runs[0].bold = None; q.runs[0].italic = italic or None
    if style: q.style = style
    return q
def table_after(p, hdr, rows, model):
    t = d.add_table(rows=1 + len(rows), cols=len(hdr)); t.style = model.style
    for j, h in enumerate(hdr): t.cell(0, j).text = h
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r): t.cell(i, j).text = str(v)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                for rr in pp.runs: rr.font.size = Pt(8.5)
    p._p.addnext(t._tbl); return t
MODEL = d.tables[16]
# Portada
p1 = P[1]; settext(p1, None, "Investment Policy Statement e Informe de Estrategia de Inversión — Versión B: la casa de Villarrica como parte de la herencia")
insert_after(P[5], "Versión paralela para discusión del grupo. Única diferencia con la versión oficial: la meta de herencia (USD 500.000 en dinero de hoy) se cumple con la cartera más la casa de Villarrica, contando solo la parte pagada con la cartera (UF 10.500) y descontando 5% de costos de venta. La asignación, las estrategias, la implementación en StockTrak y la validación son idénticas. Las cifras nuevas se calculan con el Montecarlo oficial (codigo/version_b.py) y el riesgo de precio de la casa se calibra con el Índice de Precios de Vivienda del Banco Central (Anexo E).", italic=True)
# Resumen ejecutivo
settext(find("Retorno requerido. "), "Retorno requerido. ", f"Con la casa de Villarrica como parte de la herencia, la meta financiera de herencia baja a USD {us(hf)} (USD 500.000 menos ~USD {us(cn)} de la casa, netos de costos de venta). Con los USD 144.000 invertidos y 19 aportes de USD 10.000, la cartera debe rendir {pc(tb,2)} real anual, frente a 4,65% en la versión oficial.")
settext(find("Factibilidad. "), "Factibilidad. ", f"A la tasa real libre de riesgo (2,5%), los recursos alcanzan para {pc(cb,0)} de los objetivos (cobertura de {f'{cb:.2f}'.replace('.',',')}x, frente a 0,60x en la versión oficial). El plan sigue exigiendo prima por riesgo, pero bastante menos.".replace("% de los objetivos","% de los objetivos"))
settext(find("Probabilidad de éxito. "), "Probabilidad de éxito. ", f"Educación se cumple en ~100% de los escenarios. Con retornos históricos 2013-2026, los tres objetivos se cumplen con {pc(g('Nexum','Hist','USD','P(3) B'))} ({pc(g('Nexum','Hist','UF','P(3) B'))} en UF). Con supuestos prospectivos prudentes —retorno real esperado 4,6%, ahora sobre el requerido de {pc(tb,2)}— la probabilidad conjunta es {pc(g('Nexum','Prosp','USD','P(3) B'))} ({pc(g('Nexum','Prosp','UF','P(3) B'))} en UF), frente a 46,5% en la versión oficial. La casa se valora sin plusvalía real y con la volatilidad histórica de los precios de vivienda.")
settext(find("Frente a la cartera actual de "), "Frente a la cartera actual de la AGF (85/15). ", f"La AGF rindió más en 2016-2026 (12,7% vs 9,9% anual) por su 76,5% en el S&P 500, con mayor volatilidad, caídas más profundas y ~26% de exposición a tecnología (Nexum ~14%). Con la misma regla para la casa, hacia adelante Nexum entrega mayor probabilidad de éxito: {pc(g('Nexum','Prosp','USD','P(3) B'))} vs {pc(g('AGF','Prosp','USD','P(3) B'))} en USD y {pc(g('Nexum','Prosp','UF','P(3) B'))} vs {pc(g('AGF','Prosp','UF','P(3) B'))} en UF.")
settext(find("Recomendación. "), "Recomendación. ", f"Mantener la asignación, proteger Educación como prioridad 1 y acordar con la familia que la casa de Villarrica forme parte de la herencia: sin recortar ninguna meta, la probabilidad conjunta prospectiva sube de 46,5% a {pc(g('Nexum','Prosp','USD','P(3) B'))}. Si el mercado igual rinde poco, el orden de ajuste se mantiene: primero la herencia financiera, luego Villarrica, nunca Educación.")
# 2.4
insert_after(find("La herencia tiene el mayor monto"), f"Versión B: la herencia se compone de la casa de Villarrica y de dinero. La parte de la casa pagada con la cartera (UF 10.500) vale ~USD {us(cn)} netos de 5% de costos de venta, por lo que la cartera solo debe aportar USD {us(hf)} en el año 50. Esto supone que la familia acepta que la casa, comprada para su jubilación, se herede o se venda al final del horizonte. No se usa la casa completa: las UF 15.000 de la venta de Santiago quedan fuera, como en la versión oficial.")
t2 = d.tables[2]; c = t2.rows[3].cells[1]; c.text = f"Herencia: USD 250.000 por hijo, en dinero de hoy (casa de Villarrica ~USD {us(cn)} + USD {us(hf)} de la cartera)"
# 2.5
insert_after(find("A la tasa real libre de riesgo, los recursos cubren"), f"Versión B: con la herencia financiera de USD {us(hf)}, la TIR real requerida baja de 4,65% a {pc(tb,2)} y la cobertura a la tasa libre de riesgo sube de 0,60x a {f'{cb:.2f}'.replace('.',',')}x. La Figura 1 muestra los escenarios de la versión oficial.")
p48 = find("Retorno esperado vs requerido:")
settext(p48, None, p48.text.replace("levemente bajo el requerido de 4,65%: el plan es exigente y depende de las palancas.", f"sobre el requerido de {pc(tb,2)} de la versión B (4,65% en la oficial): con la casa como parte de la herencia, el plan ya no depende de las palancas para ser factible en expectativa."))
# 4.2
p123 = find("Interpretación: Educación se cumple")
h = insert_after(p123, "Versión B: la casa de Villarrica como parte de la herencia", style=d.paragraphs[[i for i,pp in enumerate(d.paragraphs) if pp.text.startswith("Capital humano: desempleo")][0]].style)
intro = insert_after(h, f"Mismo Montecarlo (20.000 escenarios, misma semilla). En el año 50 la herencia se cumple si la cartera más la casa alcanzan USD 500.000. El precio real de la casa evoluciona entre los años 25 y 50 sin plusvalía real, con la volatilidad histórica del IPV de casas usadas del Banco Central (6,3% anual) y su correlación con la cartera; se descuenta 5% de costos de venta.")
rows = [[r.Cartera, r.Escenario, r.Moneda, pc(r["P(3) actual"]), pc(r["P(Villarrica)"]), pc(r["P(Herencia) B"]), pc(r["P(3) B"])] for _, r in R.iterrows()]
t = table_after(intro, ["Cartera", "Escenario", "Moneda", "P(3) oficial", "P(Villarrica)", "P(Herencia) B", "P(3) B"], rows, MODEL)
note = d.add_paragraph(f"Lectura: contar la casa sube la probabilidad conjunta prospectiva de 46,5% a {pc(g('Nexum','Prosp','USD','P(3) B'))} en USD. Villarrica no cambia, porque se paga antes; lo que cambia es que la herencia ya no depende solo de la cartera. La tabla de palancas que sigue corresponde a la versión oficial."); t._tbl.addnext(note._p)
# Conclusión
settext(find("La estrategia Nexum cumple el objetivo prioritario"), None, f"La estrategia Nexum cumple el objetivo prioritario —Educación— en prácticamente todos los escenarios, en USD y en UF. Si la casa de Villarrica forma parte de la herencia, el retorno requerido baja a {pc(tb,2)}, por debajo del 4,6% real esperado, y la probabilidad de cumplir las tres metas con supuestos prudentes sube de 46,5% a {pc(g('Nexum','Prosp','USD','P(3) B'))}, sin recortar ninguna meta. Frente a la cartera actual de la AGF, Nexum reduce la concentración en EE.UU. y en tecnología, baja las caídas esperadas y mejora la probabilidad de éxito. La recomendación es mantener la asignación, acordar con la familia el rol de la casa en la herencia y el orden de ajuste, y revisar anualmente el Montecarlo con los gatillos de la sección 2.9.")
# Anexo B: pregunta nueva
pq = find("Anexo C. Reproducibilidad")
ref = [pp for pp in d.paragraphs if pp.text.startswith("¿El modelo que presentan")][0]
q = insert_after(ref, ""); q.runs[0].text = "¿No es hacer trampa contar la casa en la herencia? "; q.runs[0].bold = True
q.add_run(f"No, si la familia lo acuerda: la casa es un activo que se hereda igual que la cartera. Contamos solo la parte pagada con la cartera, descontamos costos de venta y modelamos el riesgo de su precio con datos del Banco Central. Aun si la casa cae 30% al venderla, la probabilidad conjunta queda en {pc(CH[(CH.Moneda=='USD') & CH['Escenario (prospectivo)'].str.startswith('Caída permanente')]['P(3 objetivos)'].iloc[0])}, sobre el 46,5% oficial. La desventaja: la casa no es líquida y su valor depende del mercado de Villarrica.")
# Anexo E
last = d.paragraphs[-1]
hE = d.add_paragraph("Anexo E. La casa como parte de la herencia: datos y escenarios de estrés", style=d.paragraphs[[i for i,pp in enumerate(d.paragraphs) if pp.text.startswith("Anexo D.")][0]].style)
d.add_paragraph("Índice de Precios de Vivienda del Banco Central, casas usadas (trimestral, mar-2002 a mar-2026; se calcula con precios de transacciones en UF, por lo que mide precios reales): crecimiento real de 4,4% anual, volatilidad anual de 6,3%, peor caída desde un máximo de −14,2%, peor año −12,0% y correlación trimestral con la cartera Nexum de 0,04. El IPV es un índice nacional y experimental, sujeto a revisiones; una casa en Villarrica puede comportarse distinto y los índices suavizan la volatilidad de una vivienda individual, por eso se agregan escenarios de estrés.")
sp = d.add_paragraph("Escenarios (prospectivo, probabilidad de cumplir las tres metas):")
rows = [[r["Escenario (prospectivo)"], pc(r["P(3 objetivos)"]) if r.Moneda == "USD" else ""] for _, r in CH.iterrows() if r.Moneda == "USD"]
rows = [[a, b, pc(CH[(CH.Moneda == 'UF') & (CH['Escenario (prospectivo)'] == a)]['P(3 objetivos)'].iloc[0])] for a, b in rows]
table_after(sp, ["Escenario", "USD", "UF"], rows, MODEL)
d.save("../informe/Informe_Final_Nexum_VersionB.docx"); print("OK")
