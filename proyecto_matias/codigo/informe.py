"""Genera informe/Informe_Final_Nexum.docx con cifras tomadas de resultados/estado.pkl."""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from config import *

S = pd.read_pickle(os.path.join(RESULTADOS, "estado.pkl")); T = S["T"]; FIG = os.path.join(BASE, "informe", "figuras")
AZ = RGBColor(0x1F, 0x3A, 0x5F)
d = Document()
for s in d.sections: s.left_margin = s.right_margin = Cm(2.3); s.top_margin = s.bottom_margin = Cm(2.0)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.12
for h, sz in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)]:
    x = d.styles[h]; x.font.name = "Arial"; x.font.size = Pt(sz); x.font.color.rgb = AZ; x.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
p_ = lambda v, n=1: "—" if pd.isna(v) else (f"{v*100:.{n}f}%").replace(".", ",")
n_ = lambda v, n=2: "—" if pd.isna(v) else (f"{v:,.{n}f}").replace(",", "X").replace(".", ",").replace("X", ".")
u_ = lambda v: "USD " + n_(v, 0)
def H(t, l=1): d.add_heading(t, level=l)
def P(t, b=None, it=False, sz=None, al=None):
    p = d.add_paragraph()
    if b: r = p.add_run(b); r.bold = True
    r = p.add_run(t); r.italic = it
    if sz: r.font.size = Pt(sz)
    if al: p.alignment = al
    return p
def B(items):
    for x in items:
        p = d.add_paragraph(style="List Bullet")
        if isinstance(x, tuple): r = p.add_run(x[0]); r.bold = True; p.add_run(x[1])
        else: p.add_run(x)
def shade(c, col):
    pr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), col); pr.append(sh)
def TB(hdr, rows, w=None, sz=8.5, nota=None):
    t = d.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(sz); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1F3A5F")
    for k, row in enumerate(rows):
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ""; r = cs[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(sz)
            if k % 2: shade(cs[i], "F2F5F9")
    if w:
        for row in t.rows:
            for i, x in enumerate(w): row.cells[i].width = Cm(x)
    if nota: P(nota, it=True, sz=8)
    else: d.add_paragraph()
def FIGURA(n, cap, w=16):
    d.add_picture(os.path.join(FIG, n), width=Cm(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; P(cap, it=True, sz=8, al=WD_ALIGN_PARAGRAPH.CENTER)

mc = T["Montecarlo"]; g = lambda c, e, m, k: mc[(mc.Cartera == c) & mc.Escenario.str.startswith(e) & (mc.Moneda == m)][k].iloc[0]
fac = S["fac"]; met = T["Metricas"]; mt = lambda per, s, k: met[(met["Período"] == per) & (met.Serie == s)][k].iloc[0]
te = T["Tracking_error"].set_index("Concepto")["Valor"]; tir = T["TIR"].set_index("Escenario")["TIR real requerida"]
pal = T["Palancas"].set_index("Palanca (prospectivo, USD)"); ic = T["Impacto_correcciones"]; asg = T["Asignacion"]; av = T["Aversion_riesgo"]["Valor"].tolist()
wf = T["Walk_forward"]; wfv = lambda c, per: wf[(wf.Cartera == c) & wf["Período"].str.startswith(per)]["Sharpe"].iloc[0]
rb = T["Robustez"]; rbv = lambda e, v: rb[(rb.Estrategia == e) & rb.Variante.str.startswith(v)]["Sharpe"].iloc[0]
es = T["Estrategias"]; esv = lambda e, k: es[es.Estrategia == e][k].iloc[0]

# ================= PORTADA =================
P("Nexum Financial Advisors", al=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(22)
d.paragraphs[-1].runs[0].bold = True; d.paragraphs[-1].runs[0].font.color.rgb = AZ
P("Investment Policy Statement e Informe de Estrategia de Inversión", al=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(14)
P("Familia Allende Cunich · Mandato USD 150.000 + aportes anuales", al=WD_ALIGN_PARAGRAPH.CENTER)
P("Grupo 11 — M. José Iannello, Martín Lillo, Francisca Morales, Vicente Salah, Matías Torres", it=True, al=WD_ALIGN_PARAGRAPH.CENTER)
P("ESF 2026-2 · Universidad Adolfo Ibáñez · Datos de mercado al 31-ago-2026; implementación al 5-oct-2026", it=True, sz=9, al=WD_ALIGN_PARAGRAPH.CENTER)
P("Todas las cifras de este documento son reproducibles con codigo/run_all.py y coinciden con modelo/Modelo_Nexum_v2.xlsx.", it=True, sz=8.5, al=WD_ALIGN_PARAGRAPH.CENTER)
d.add_page_break()

# ================= 1. RESUMEN EJECUTIVO =================
H("1. Resumen ejecutivo")
P("¿Se cumplen los objetivos de la familia y con qué riesgo? ", b="La pregunta del cliente: ")
B([("Retorno requerido. ", f"Para financiar Educación, Villarrica y la herencia con USD 150.000 y 19 aportes de USD 10.000, la cartera debe rendir {p_(fac['TIR requerida'],2)} real anual (TIR de los flujos del enunciado). Con supuestos más exigentes y fieles al caso —aportes en USD nominales, arancel creciendo 2% real, 14 años efectivos de aportes— el requerimiento sube a {p_(tir.iloc[5],1)}."),
   ("Factibilidad. ", f"A la tasa real libre de riesgo (2,5%), los objetivos valen {u_(fac['VP objetivos'])} y los recursos {u_(fac['VP recursos (capital + aportes)'])}: cobertura de {n_(fac['Cobertura (recursos / objetivos)'],2)}x. El plan exige capturar prima por riesgo. Villarrica explica {p_(fac['Peso VP Villarrica'],0)} del valor presente de los objetivos."),
   ("Asignación. ", "Núcleo pasivo 69,4% (S&P 500 37,9%, Tesoro 10-20 años 16,9%, corporativos largos 14,6%) y satélite activo 30,6% (Low-Volatility 17,5%, Momentum en acciones 7,1%, Momentum en commodities 6,1%), obtenido con Black-Litterman. Renta variable efectiva 62%."),
   ("Probabilidad de éxito. ", f"Educación se cumple en ~100% de los escenarios simulados, en USD y en UF. Con retornos históricos 2013-2026, los tres objetivos se cumplen con {p_(g('Nexum','Hist','USD','P(3 objetivos)'))} ({p_(g('Nexum','Hist','UF','P(3 objetivos)'))} medido en UF). Con supuestos prospectivos prudentes —retorno real esperado {p_(g('Nexum','Prosp','USD','Retorno real anual'))}, prácticamente igual al requerido— la probabilidad conjunta es {p_(g('Nexum','Prosp','USD','P(3 objetivos)'))}."),
   ("Frente a la cartera actual de la AGF (85/15). ", f"La AGF rindió más en 2016-2026 ({p_(mt('2016-2026','Cartera actual AGF 85/15','CAGR'))} vs {p_(mt('2016-2026','Cartera Nexum','CAGR'))} anual) por su 76,5% en el S&P 500, con mayor volatilidad, caídas más profundas y ~26% de exposición a tecnología (el mismo riesgo del empleo de Tomás; Nexum ~14%). Hacia adelante, Nexum entrega mayor probabilidad de éxito ({p_(g('Nexum','Prosp','UF','P(3 objetivos)'))} vs {p_(g('AGF 85/15','Prosp','UF','P(3 objetivos)'))} en UF)."),
   ("Gestión activa. ", f"Tracking error ex-ante {p_(te.iloc[0])} (objetivo 3%, máximo 4%) e information ratio {n_(mt('2016-2026','Cartera Nexum','Information ratio'))} frente al benchmark de política en 2016-2026."),
   ("Recomendación. ", f"Mantener la asignación, proteger Educación como prioridad 1 y acordar con la familia el orden de ajuste si el mercado no entrega el retorno requerido: primero la herencia, luego Villarrica, nunca Educación. Reducir Villarrica 25% y la herencia a USD 250.000 eleva la probabilidad conjunta prospectiva de {p_(pal.loc['Base','P(3 objetivos)'])} a {p_(pal.loc['Ambas','P(3 objetivos)'])}.")])

# ================= 2. IPS =================
d.add_page_break(); H("2. Investment Policy Statement")
H("2.1 Propósito y alcance", 2)
P("Este IPS es el acuerdo entre Nexum Financial Advisors (el Asesor) y María Ignacia Cunich y Tomás Allende (el Cliente) para administrar USD 150.000 hoy invertidos en fondos mutuos de una AGF chilena, más aportes de USD 10.000 cada abril hasta que Tomás cumpla 65 años. Define objetivos, restricciones, asignación de activos, reglas de selección y rebalanceo, métricas de evaluación y la política de revisión. Toda decisión de inversión se toma dentro de este marco.")
H("2.2 Deberes y responsabilidades", 2)
P("El Asesor actúa como fiduciario, bajo el Código de Ética y los Estándares de Conducta Profesional del CFA Institute y la normativa de la CMF. Sus responsabilidades: recomendar una asignación coherente con el perfil, seleccionar los instrumentos, ejecutar y rebalancear, monitorear el desempeño contra el benchmark e informar al Cliente. El Cliente debe informar oportunamente cualquier cambio material (empleo, patrimonio, objetivos), lo que gatilla la revisión de este IPS.")
TB(["Control de cumplimiento (CMF / Ley 19.913)", "Meta medible"], [
    ["Conozca a su cliente: identidad, condición de Persona Expuesta Políticamente y perfil transaccional", "100% antes de operar"],
    ["Origen de fondos: acreditar que los USD 150.000 provienen de ahorro laboral en una AGF regulada; documentar cada aporte anual", "100% de los aportes"],
    ["Monitoreo de operaciones inusuales respecto del perfil (montos o movimientos no explicados)", "Revisión mensual"],
    ["Reporte de Operación Sospechosa a la UAF, con deber de reserva", "Dentro de 48 horas desde la detección"],
    ["Actualización del perfil e idoneidad de los instrumentos", "Anual o ante un cambio material"],
    ["Conservación de registros", "Mínimo 5 años"]], w=[11.5, 4.5])
H("2.3 Perfil del inversionista", 2)
P("María Ignacia (40 años, abogada en consumo masivo) y Tomás (46, ingeniero en software) tienen dos hijos (Vicente, 4; Emilia, 2), ingresos estables y educación financiera media, sin experiencia en derivados ni criptoactivos. El perfil distingue dos dimensiones:")
TB(["Dimensión", "Evidencia", "Conclusión"], [
    ["Capacidad de riesgo (objetiva)", f"Capital USD 150.000 + valor presente de los aportes (~USD 150.000 a 2,5% real); sin necesidades de liquidez hasta el año 14; horizonte hasta 50 años", "Alta en el largo plazo; decreciente frente a Educación (años 14-20)"],
    ["Disposición al riesgo (conductual)", "Toleran hoy 85% en activos de riesgo (cartera AGF) y piden activamente commodities y alternativos", "Moderada-agresiva; renta variable de 60-65%"],
    ["Capital humano", "El ingreso de Tomás depende del sector software (se comporta como una acción tecnológica); el de María Ignacia, de consumo masivo (como un bono)", "Limitar tecnología: ~14% en Nexum vs ~26% en la AGF"],
    ["Aversión al riesgo (A)", f"A del inversionista de mercado (S&P 500, 2013-2026) = {n_(av[0],1)}; A implícito en la cartera recomendada = {n_(av[1],1)}; δ de Black-Litterman = {n_(av[2],1)}", "El cliente acepta más riesgo que el inversionista promedio, coherente con su horizonte"]], w=[3.4, 7.6, 5])
P("El cuestionario de tolerancia al riesgo (Grable y Lytton) se aplicará directamente a ambos clientes en la reunión de inicio; si su resultado difiere de este perfil, se revisará el IPS. No se usan respuestas simuladas.", it=True, sz=9)
H("2.4 Objetivos, prioridad y horizontes", 2)
pd_ = T["Plan_desembolsos"]
TB(["Prioridad", "Objetivo (moneda del cliente)", "Años", "UF", "USD reales", "Flexibilidad"], [
    ["1", "Educación: CLP 65 millones por hijo, en 5 pagos anuales", "14-18 y 16-20", n_(pd_[pd_.Concepto == 'Educación']['UF'].sum(), 0), n_(pd_[pd_.Concepto == 'Educación']['USD reales'].sum(), 0), "Rígida, no diferible"],
    ["2", "Villarrica: desembolso UF 10.500 (UF 15.000 adicionales vienen de la venta de la casa, fuera de la cartera)", "25", "10.500", n_(VILLARRICA_UF * UF_CLP / USD_CLP, 0), "Postergable o reducible"],
    ["3", "Herencia: USD 250.000 por hijo, en dinero de hoy", "50", n_(HERENCIA_USD * USD_CLP / UF_CLP, 0), "500.000", "Residual: absorbe la variabilidad"]], w=[1.5, 6.4, 1.9, 1.6, 1.9, 2.7])
P("La herencia tiene el mayor monto pero la menor prioridad: no tiene fecha rígida y es la variable de ajuste natural. Villarrica tiene fecha, pero el propio enunciado admite postergación, reducción de superficie o una localización alternativa. Educación no admite sustitución. Horizontes: corto plazo (menos de 3 años antes de cada pago de Educación), mediano plazo (3 a 15 años) y largo plazo (Villarrica y herencia). Entre los años 16 y 18 se pagan las dos carreras a la vez: es la mayor carga de liquidez del plan.")
H("2.5 Retorno requerido y factibilidad", 2)
FIGURA("f5_tir.png", "Figura 1. TIR real requerida por escenario. Fuente: codigo/objetivos.py; hoja 02_Objetivos_TIR (fórmula IRR).")
P(f"A la tasa real libre de riesgo, los recursos cubren {n_(fac['Cobertura (recursos / objetivos)'],2)} veces los objetivos: el plan es factible solo si la cartera captura prima por riesgo. Valorizados a la TIR, Villarrica representa {p_(fac['Peso VP Villarrica'],0)} de la exigencia, Educación {p_(fac['Peso VP Educación'],0)} y la herencia {p_(fac['Peso VP Herencia'],0)}. No se usa la venta de la casa de Santiago más allá de las UF 15.000 que el enunciado asigna a Villarrica.")
TB(["Año", "Concepto", "UF", "USD reales"], [[int(r["Año"]), r["Concepto"][:80], n_(r["UF"], 1), n_(r["USD reales"], 0)] for _, r in pd_.iterrows()], w=[1.3, 10, 2.2, 2.5], nota="Plan de desembolsos. Tipo de cambio de corte: UF = CLP 40.846; USD = CLP 911,77 (BCCh, 11-ago-2026).")
H("2.6 Restricciones", 2)
B([("Liquidez: ", "sin retiros hasta el año 14; reserva operativa de 4% del capital (2% caja + 2% reserva cambiaria); fondo de emergencia fuera del mandato."),
   ("Legales y tributarias: ", "ETF y acciones listadas en EE.UU.; retención de 15% sobre dividendos por el convenio Chile-EE.UU.; tributación en Chile al liquidar."),
   ("Éticas e idoneidad (CFA, Estándar III.C): ", "sin derivados complejos, criptoactivos, apalancamiento ni ventas cortas, por el nivel de conocimiento declarado."),
   ("Plataforma: ", "instrumentos disponibles en StockTrak; al menos 150 operaciones en 12 semanas y 5 por semana en promedio.")])
H("2.7 Asignación estratégica y riesgo-retorno esperado", 2)
FIGURA("f3_asignacion.png", "Figura 2. Pesos Black-Litterman (aprobados) frente al prior de paridad de riesgo y a Markowitz histórico. Fuente: hojas 05 y 06 del modelo.")
P("Marco: Black-Litterman parte de un prior de equilibrio de paridad de riesgo (peso inverso a la volatilidad), deriva retornos implícitos π = δΣw y los combina con dos vistas explícitas: el S&P 500 rinde 6% (prima por riesgo de largo plazo) y Low-Volatility supera a commodities en 2 puntos. La optimización restringe el núcleo pasivo entre 60% y 70%. Markowitz con medias históricas se reporta solo como referencia: concentra en lo que más rindió en el pasado (ver walk-forward, sección 4.4).")
P(f"Los pesos aprobados en septiembre siguen vigentes: recalibrados con los datos corregidos de esta versión, difieren en menos de {n_(abs(asg['Diferencia (pp)']).max(),1)} puntos, dentro de la banda de rebalanceo de ±2 puntos. Cambiarlos solo generaría costos.")
m16 = lambda s, k: mt('2016-2026', s, k)
TB(["Métrica (2016-2026, USD)", "Nexum", "Benchmark", "AGF 85/15", "S&P 500"],
   [[k] + [(p_(m16(s, k)) if k not in ("Sharpe", "Sortino", "Beta", "Information ratio") else n_(m16(s, k))) for s in ["Cartera Nexum", "Benchmark de política", "Cartera actual AGF 85/15", "S&P 500 (SPY)"]]
    for k in ["CAGR", "Volatilidad", "Sharpe", "Sortino", "Beta", "Máx. caída", "VaR 95% mensual", "CVaR 95% mensual", "Tracking error", "Information ratio"]], w=[5, 2.7, 2.7, 2.7, 2.7])
P(f"Retorno esperado vs requerido: el retorno histórico real 2013-2026 fue {p_(g('Nexum','Hist','USD','Retorno real anual'))}, muy sobre el requerido. Con supuestos prospectivos (tasa real ~2,5% más primas moderadas; sección 4.3) el retorno real esperado es {p_(g('Nexum','Prosp','USD','Retorno real anual'))}, igual al requerido de {p_(fac['TIR requerida'],2)}. Nivel de riesgo aceptable: volatilidad anual de 10-12%, máxima caída tolerada de 25% y probabilidad de Educación de al menos 90%.")
FIGURA("f1_crecimiento.png", "Figura 3. Crecimiento de USD 100 (2016-2026). Fuente: resultados/retornos_cartera_benchmark.csv.")
FIGURA("f2_caidas.png", "Figura 4. Caídas desde el máximo: Nexum frente a la cartera actual de la AGF.")
H("2.8 Selección, seguimiento y rebalanceo", 2)
TB(["Elemento", "Regla"], [
    ["Selección", "Núcleo: ETF líquidos de bajo costo. Satélite: acciones del S&P 500 (precio ≥ USD 5, liquidez ≥ USD 5 millones/día) y 5 ETF de commodities"],
    ["Señales", "Mensuales, con el cierre del último día hábil (codigo/senales.py); ejecución escalonada durante la semana siguiente"],
    ["Bandas", "Núcleo ±5 puntos por activo; satélite ±2 puntos por sleeve; dentro de cada sleeve, peso igual"],
    ["Revisión semanal", "Ranking de Momentum con banda de permanencia: una acción se mantiene mientras esté en el top 25"],
    ["Costos", "No se ejecutan ajustes menores a USD 500 (comisión de USD 10 = 2% del monto)"],
    ["Disciplina de venta", "Núcleo sin stop-loss (se controla con bandas). Satélite: salida por señal; stop de emergencia de −15%; en posiciones con ganancia, stop móvil sobre el costo"],
    ["Eventos", "No se ejecuta en la primera media hora ni durante la publicación de datos de la Fed, inflación o empleo"]], w=[3.6, 12.4])
H("2.9 Política de revisión y control", 2)
B(["Revisión anual del IPS y del Montecarlo con datos actualizados.",
   "Revisión extraordinaria si: la probabilidad de Educación cae bajo 90%; la cartera cae más de 15% desde su máximo; el tracking error supera 4%; Tomás pierde su empleo; cambia algún objetivo.",
   "Informe mensual al cliente (desempeño frente al benchmark, operaciones y riesgos) y reunión trimestral, donde los clientes participan en las decisiones tácticas."])

# ================= 3. ESTRATEGIA =================
d.add_page_break(); H("3. Estrategia de inversión")
H("3.1 Arquitectura núcleo-satélite", 2)
P("El núcleo pasivo (69%) captura la prima de mercado y de plazo a bajo costo y es la base de financiamiento de largo plazo. El satélite (31%) busca primas de factores documentadas, con un presupuesto de riesgo activo acotado. Las tres estrategias se eligieron porque se complementan: Momentum aporta retorno en mercados con tendencia, Low-Volatility reduce la beta de la cartera y Momentum en commodities diversifica frente a acciones y bonos.")
H("3.2 Estrategias activas: evidencia y resultados", 2)
TB(["Estrategia", "Implementación", "Evidencia académica", "Riesgo/retorno teórico"], [
    ["Momentum en acciones", "Top 10 del S&P 500 por 50% retorno 6m + 50% retorno 12m; filtros de tendencia y liquidez; máx. 25% por sector; mensual",
     "Jegadeesh y Titman (1993); Carhart (1997); Asness, Moskowitz y Pedersen (2013)", "Ganadores menos perdedores ~1% mensual (J&T, 1965-1989); caídas abruptas en rebotes del mercado (Daniel y Moskowitz, 2016)"],
    ["Low-Volatility", "Top 10 del S&P 500 por menor volatilidad 12m; peso igual; máx. 25% por sector; solo compras",
     "Haugen y Heins (1975); Ang et al. (2006); Blitz y van Vliet (2007); Baker, Bradley y Wurgler (2011)", "Retorno similar al mercado con ~30% menos volatilidad; beta 0,5-0,7"],
    ["Momentum en commodities", "GLD, USO, DBA, SLV, CPER: compra los que superan filtros de tendencia 3/6/12 meses y media de 200 días, con vol. < 50%",
     "Erb y Harvey (2006); Miffre y Rallis (2007); Moskowitz, Ooi y Pedersen (2012)", "Prima del orden de 9% anual en futuros (Miffre y Rallis); baja correlación con acciones y bonos"]], w=[2.8, 5, 4.2, 4.2], sz=8)
TB(["Estrategia (2016-2026)", "Benchmark", "Retorno anual", "Volatilidad", "Sharpe", "Beta", "Máx. caída", "TE", "IR"],
   [[NOM, BEN, p_(esv(NOM, "CAGR")), p_(esv(NOM, "Volatilidad")), n_(esv(NOM, "Sharpe")), n_(esv(NOM, "Beta")), p_(esv(NOM, "Máx. caída")), p_(esv(NOM, "Tracking error")), n_(esv(NOM, "Information ratio"))]
    for NOM, BEN in [("Momentum en acciones", "QMOM"), ("Low-Volatility", "SPLV"), ("Momentum en commodities", "Mezcla 5 ETF")]], w=[3.6, 2.2, 1.7, 1.7, 1.3, 1.2, 1.7, 1.3, 1.3], sz=8)
P("Fundamento de los filtros: 6 y 12 meses son las ventanas de formación estándar (Jegadeesh y Titman, 1993); la media de 200 días es un filtro de tendencia clásico (Faber, 2007); el filtro de Sharpe descarta alzas con volatilidad excesiva, en la línea del control de volatilidad de Barroso y Santa-Clara (2015); el límite sectorial controla la concentración en tecnología, que además es el riesgo del capital humano de Tomás.")
P("Lo que se implementa es Low-Volatility y no Betting Against Beta: la estrategia de Frazzini y Pedersen (2014) apalanca las acciones de beta baja y vende en corto las de beta alta. Nexum no usa apalancamiento ni ventas cortas, así que su retorno proviene de la anomalía de baja volatilidad y conserva una beta de ~0,55.")
P(f"Discusión crítica. (i) Las primas de factores se reducen fuera de muestra y después de su publicación (McLean y Pontiff, 2016), por eso los supuestos prospectivos usan primas menores a las históricas. (ii) Momentum sufre caídas abruptas cuando el mercado rebota. (iii) Low-Volatility se concentra en utilities y REITs, sensibles a tasas; en septiembre de 2026 se observó en vivo, con la tasa a 10 años en 5,3%. (iv) Momentum en commodities es sensible a sus umbrales y su universo se amplió mirando el backtest (sección 4.4): con el universo original de 3 ETF su Sharpe cae de {n_(rbv('Commodities','BASE'))} a {n_(rbv('Commodities','Universo original'))}. Por eso es el sleeve con menor peso.")
H("3.3 Alternativas evaluadas y descartadas", 2)
B(["Rotación sectorial con ETF: misma familia que Momentum (correlación ~0,67) y no mejora el Sharpe de la cartera; se pierde el rol defensivo de Low-Vol.",
   "Reversión de corto plazo semanal: ~19 operaciones por semana; los costos (~10% anual a 0,1% por lado y peor con la comisión fija de StockTrak) eliminan su prima.",
   "Momentum de 4 semanas y estrategias semanales: Sharpe neto de costos de 0,2 a 0,5, inferior a las estrategias mensuales.",
   "Betting Against Beta: requiere apalancamiento y ventas cortas, excluidos por idoneidad."])
H("3.4 Moneda funcional: USD, con simulaciones en UF", 2)
fx = T["Riesgo_cambiario"].set_index("Indicador")["Valor"]
P("La moneda funcional es el USD porque todo el universo de inversión cotiza en dólares y cubrir el tipo de cambio exigiría derivados que el perfil excluye. Las necesidades del cliente (UF/CLP) se evalúan explícitamente: el Montecarlo se presenta también en UF.")
B([("Cobertura natural: ", f"el peso se deprecia cuando caen los mercados globales (correlación S&P 500 vs. USD/CLP de {n_(fx.iloc[1])}). En el 10% de los peores meses del S&P 500 el dólar subió {p_(fx.iloc[2])} en promedio: la cartera perdió {p_(fx.iloc[3])} en USD pero solo {p_(fx.iloc[4])} medida en pesos."),
   ("Diversificación del riesgo país: ", "el capital humano y la vivienda de la familia ya están expuestos a Chile."),
   ("Costo: ", f"la volatilidad del USD/CLP es {p_(fx.iloc[0])} anual; medida en pesos, la volatilidad de la cartera sube de {p_(fx.iloc[5])} a {p_(fx.iloc[6])}, y la probabilidad conjunta histórica baja de {p_(g('Nexum','Hist','USD','P(3 objetivos)'))} a {p_(g('Nexum','Hist','UF','P(3 objetivos)'))}. Se monitorea en cada revisión.")])
H("3.5 Benchmark de política", 2)
TB(["Sleeve", "Peso", "Benchmark (índice invertible del mismo universo)"], [
    ["S&P 500", "37,9%", "S&P 500 (SPY)"], ["Tesoro 10-20 años", "16,9%", "ICE U.S. Treasury 10-20 Year Index (TLH)"], ["Corporativos largo plazo", "14,6%", "Bloomberg U.S. Long Corporate Index (VCLT)"],
    ["Momentum en acciones", "7,1%", "Alpha Architect U.S. Quantitative Momentum (QMOM)"], ["Low-Volatility", "17,5%", "S&P 500 Low Volatility Index (SPLV)"],
    ["Momentum en commodities", "6,1%", "Mezcla igual ponderada GLD/USO/DBA/SLV/CPER"]], w=[4.2, 1.6, 10.2],
   nota="Benchmark compuesto ponderado por la asignación estratégica y expresado en USD. Reemplaza la tasa de política de la Fed y el MSCI ACWI de la versión anterior, que no eran invertibles ni representaban el universo. Limitación: QMOM existe desde dic-2015, por lo que las métricas relativas usan 2016-2026.")
H("3.6 Métricas de desempeño", 2)
TB(["Métrica", "Para qué se usa", "Limitación"], [
    ["Sharpe", "Retorno por unidad de riesgo total", "Supone normalidad; penaliza la volatilidad al alza"],
    ["Sortino", "Penaliza solo la volatilidad a la baja", "Inestable con pocas observaciones"],
    ["Treynor y alfa de Jensen", "Retorno por unidad de riesgo sistemático y valor agregado ajustado por beta", "Dependen del modelo de mercado y del beta estimado"],
    ["Tracking error e information ratio", "Riesgo activo y su eficiencia frente al benchmark", "Requieren un benchmark bien construido"],
    ["Máxima caída y Calmar", "Tolerancia conductual a pérdidas", "Dependen de la ventana"],
    ["VaR y CVaR 95% mensual", "Pérdida en un mes malo y en la cola", "El VaR no informa el tamaño de la cola; por eso se agrega CVaR"],
    ["Probabilidad de éxito por objetivo", "Métrica central de un IPS por objetivos", "Depende de los supuestos del Montecarlo"]], w=[4, 6.2, 5.8])
H("3.7 Tracking error objetivo", 2)
P(f"El límite se deriva del presupuesto de riesgo activo (Grinold y Kahn, 2000): con un alfa objetivo de 1,5% anual y un information ratio objetivo de 0,5, el TE objetivo es 1,5% / 0,5 = 3,0%; el máximo es 4% (alfa de 2% con el mismo IR). El TE ex-ante con los pesos aprobados es {p_(te.iloc[0])} (2016-2026) y {p_(te.iloc[1])} en los últimos 36 meses, dentro del objetivo; el information ratio realizado es {n_(mt('2016-2026','Cartera Nexum','Information ratio'))}.")

# ================= 4. BACKTEST Y MONTECARLO =================
d.add_page_break(); H("4. Backtesting y simulación de Montecarlo")
H("4.1 Datos, período y supuestos", 2)
B(["Acciones del S&P 500 con membresía histórica point-in-time (2009-2026): sin sesgo de supervivencia en los miembros del índice.",
   "ETF del núcleo, de commodities y benchmarks: Yahoo Finance, precios ajustados por dividendos.",
   "Ventana de evaluación: ene-2013 a ago-2026 (164 meses); métricas relativas, 2016-2026 (128 meses).",
   "Señal con información disponible al cierre del mes t; retorno desde t+1 (sin anticipación). Costo de 0,1% sobre la rotación.",
   "Montecarlo: 20.000 escenarios de 50 años por bootstrap mensual; inflación de EE.UU. de 2%; capital invertido de USD 144.000 (neto de la reserva de 4%)."])
H("4.2 Montecarlo: probabilidad de cumplir los objetivos", 2)
FIGURA("f4_montecarlo.png", "Figura 5. Probabilidad de cumplir cada objetivo. Fuente: codigo/montecarlo.py; hoja 08_Montecarlo.")
TB(["Cartera", "Escenario", "Moneda", "Retorno real", "P(Educación)", "P(Villarrica)", "P(Herencia)", "P(3 objetivos)"],
   [[r.Cartera, r.Escenario, r.Moneda, p_(r["Retorno real anual"]), p_(r["P(Educación)"]), p_(r["P(Villarrica)"]), p_(r["P(Herencia)"]), p_(r["P(3 objetivos)"])] for _, r in mc.iterrows()],
   w=[2, 3.2, 1.4, 1.7, 1.9, 1.9, 1.9, 2], sz=8,
   nota="Escenario prospectivo: misma volatilidad y correlaciones históricas, con medias ajustadas a supuestos de mercado (retornos reales: S&P 500 5,0%, Low-Vol 4,5%, Momentum 5,5%, TLH 2,6%, VCLT 3,2%, commodities 2,0%). UF: tipo de cambio real sin tendencia, con la volatilidad y la correlación observadas.")
P("Interpretación: Educación se cumple en prácticamente todos los escenarios, por lo que el piso de diseño (90-95%) se satisface con holgura. La incertidumbre se concentra en Villarrica y la herencia, que el IPS define como flexibles. Con retornos históricos el plan es holgado; con supuestos prudentes la probabilidad conjunta es menor a 50%, coherente con un retorno esperado que apenas iguala al requerido. Esta es la conclusión que se presenta a los clientes, junto con las palancas de ajuste:")
TB(["Palanca (escenario prospectivo, USD)", "P(Educación)", "P(Villarrica)", "P(Herencia)", "P(3 objetivos)"], [[k] + [p_(v) for v in r.values] for k, r in pal.iterrows()], w=[6.4, 2.4, 2.4, 2.4, 2.4])
H("4.3 Validación del modelo", 2)
P("Correcciones respecto de la versión anterior del modelo. La auditoría encontró tres problemas que afectaban los resultados; esta versión los corrige y mide su efecto:")
TB(["Versión", "Meses", "Momentum", "Low-Vol", "Commod.", "Cartera CAGR", "Sharpe", "BL SPY", "BL LV"],
   [[r["Versión"][:52], int(r["Meses"]), p_(r["CAGR Momentum"]), p_(r["CAGR Low-Vol"]), p_(r["CAGR Commodities"]), p_(r["Cartera (pesos aprobados) CAGR"]), n_(r["Cartera Sharpe"]), p_(r["BL SPY"]), p_(r["BL LOW_VOL"])] for _, r in ic.iterrows()],
   w=[5.6, 1.1, 1.5, 1.4, 1.4, 1.6, 1.2, 1.3, 1.3], sz=7.5,
   nota="(1) Low-Vol devolvía datos vacíos en jul-2018 y dic-2018 y el modelo eliminaba esos meses de las 6 series, incluido el peor mes del período para las acciones. (2) El filtro de Sharpe de Momentum usaba la volatilidad de un mes en el backtest y de 12 meses en la operación. (3) Sensibilidad: contar el costo en compras y ventas. Ninguna corrección mueve los pesos más de 0,7 puntos.")
FIGURA("f6_robustez.png", "Figura 6. Robustez: Sharpe 2013-2026 de cada variante de parámetros (un cambio a la vez). Fuente: codigo/validacion.py.")
P(f"Robustez. Momentum y Low-Volatility son estables frente a sus parámetros (Sharpe entre {n_(rb[rb.Estrategia!='Commodities'].Sharpe.min())} y {n_(rb[rb.Estrategia!='Commodities'].Sharpe.max())}) y la regla vigente no es la mejor variante, lo que es evidencia en contra de sobreajuste. Momentum en commodities sí depende de sus umbrales y de la ampliación del universo de 3 a 5 ETF, que se decidió observando el backtest: se declara como sesgo de selección y explica su menor peso.")
TB(["Cartera", "Sharpe 2013-2019 (dentro de muestra)", "Sharpe 2020-2026 (fuera de muestra)"],
   [[c, n_(wfv(c, "Dentro")), n_(wfv(c, "Fuera"))] for c in wf.Cartera.unique()], w=[7.4, 4.3, 4.3],
   nota="Walk-forward: los pesos se calibran solo con 2013-2019 y se evalúan en 2020-2026. Markowitz con medias históricas sobreajusta; Black-Litterman calibrado con la misma información mantiene el mejor desempeño fuera de muestra, lo que respalda su uso.")
P("Sesgos y limitaciones declarados:", b=None)
B(["Clasificación sectorial actual (no histórica) en el límite sectorial: sesgo de anticipación leve; afecta sobre todo antes de 2016 y a ~3% de posiciones sin clasificar.",
   "Precio faltante en el mes: se trata como 0% en esa posición.",
   "Commodities: universo ampliado mirando el backtest (sesgo de selección), controlado con el universo original.",
   "Montecarlo: bootstrap independiente (no captura regímenes ni autocorrelación); los supuestos prospectivos son propios, no pronósticos."])

# ================= 5. IMPLEMENTACIÓN =================
d.add_page_break(); H("5. Implementación en StockTrak (al 5-oct-2026)")
P("La cartera se implementó el 14-sep-2026 con los pesos aprobados. Decisiones relevantes durante la ejecución:")
B([("Rebalanceo de fin de septiembre: ", "Momentum se llevó a peso igual; Moderna se redujo de 12 a 5 acciones, asegurando ~USD 225 de ganancia, con un stop móvil sobre el costo para las restantes."),
   ("Low-Volatility: ", "se redujo utilities y real estate de 50% a ~29% del sleeve (salieron Realty Income y Duke Energy), por la sensibilidad a tasas observada; entraron PG, RSG y SNA siguiendo el ranking."),
   ("Petróleo: ", "se vendieron MPC, VLO y USO para reducir la exposición conjunta al petróleo; la señal de octubre confirmó la salida de USO."),
   ("Operaciones: ", "~50 en las primeras tres semanas, en línea con la regla de 150 en 12 semanas.")])
P("Momentum es el sleeve con mejor desempeño en vivo (≈ +3,7% desde el 14-sep frente a ≈ +2,4% de QMOM). La cartera total se ve afectada por la subida de la tasa a 10 años a 5,3%, que golpea al núcleo de renta fija larga (~30% de la cartera). Es el principal riesgo de corto plazo y está identificado en la sección 6.", it=False)

# ================= 6. RIESGOS =================
H("6. Riesgos principales y mitigantes")
TB(["Riesgo", "Exposición", "Mitigante"], [
    ["Tasas de interés", "TLH + VCLT ≈ 31%; más Low-Vol en utilities/REITs", "Banda de rebalanceo; límite a utilities/REITs; revisión de la duración en la revisión anual"],
    ["Retorno de mercado menor al requerido", "Retorno esperado ≈ requerido", "Palancas acordadas: herencia y luego Villarrica; nunca Educación"],
    ["Concentración tecnológica / capital humano", "Empleo de Tomás en software", "Límite de 2 acciones por sector en el satélite; tecnología ~14% vs ~26% en la AGF"],
    ["Caídas abruptas de Momentum", "7% de la cartera, volatilidad ~24%", "Peso acotado, peso igual, stops de emergencia"],
    ["Tipo de cambio", "Objetivos en UF/CLP, activos en USD", "Cobertura natural (correlación negativa); simulaciones en UF; reserva cambiaria de 2%"],
    ["Empleo de Tomás", "Aportes condicionados a seguir empleado", "Escenario de 14 aportes en la TIR; fondo de emergencia fuera del mandato; seguro de vida e invalidez"]], w=[3.6, 5.6, 6.8])

# ================= 7. CONCLUSIÓN =================
H("7. Conclusión")
P(f"La estrategia Nexum cumple el objetivo prioritario —Educación— en prácticamente todos los escenarios, en USD y en UF. Villarrica y la herencia dependen del retorno que entregue el mercado: se cumplen holgadamente con retornos como los de 2013-2026 y con {p_(g('Nexum','Prosp','USD','P(3 objetivos)'),0)} de probabilidad conjunta bajo supuestos prudentes. Frente a la cartera actual de la AGF, Nexum reduce la concentración en EE.UU. y en tecnología, baja las caídas esperadas y mejora la probabilidad de éxito hacia adelante. La recomendación es mantener la asignación, revisar anualmente el Montecarlo con los gatillos de la sección 2.9 y acordar desde ya el orden de ajuste con la familia.")

# ================= REFERENCIAS =================
H("Referencias")
for ref in ["Ang, A., Hodrick, R., Xing, Y. y Zhang, X. (2006). The Cross-Section of Volatility and Expected Returns. Journal of Finance, 61(1), 259-299.",
 "Asness, C., Moskowitz, T. y Pedersen, L. (2013). Value and Momentum Everywhere. Journal of Finance, 68(3), 929-985.",
 "Baker, M., Bradley, B. y Wurgler, J. (2011). Benchmarks as Limits to Arbitrage: Understanding the Low-Volatility Anomaly. Financial Analysts Journal, 67(1), 40-54.",
 "Barroso, P. y Santa-Clara, P. (2015). Momentum Has Its Moments. Journal of Financial Economics, 116(1), 111-120.",
 "Black, F. y Litterman, R. (1992). Global Portfolio Optimization. Financial Analysts Journal, 48(5), 28-43.",
 "Blitz, D. y van Vliet, P. (2007). The Volatility Effect. Journal of Portfolio Management, 34(1), 102-113.",
 "Carhart, M. (1997). On Persistence in Mutual Fund Performance. Journal of Finance, 52(1), 57-82.",
 "Daniel, K. y Moskowitz, T. (2016). Momentum Crashes. Journal of Financial Economics, 122(2), 221-247.",
 "DeMiguel, V., Garlappi, L. y Uppal, R. (2009). Optimal Versus Naive Diversification. Review of Financial Studies, 22(5), 1915-1953.",
 "Erb, C. y Harvey, C. (2006). The Strategic and Tactical Value of Commodity Futures. Financial Analysts Journal, 62(2), 69-97.",
 "Faber, M. (2007). A Quantitative Approach to Tactical Asset Allocation. Journal of Wealth Management, 9(4), 69-79.",
 "Frazzini, A. y Pedersen, L. (2014). Betting Against Beta. Journal of Financial Economics, 111(1), 1-25.",
 "Grinold, R. y Kahn, R. (2000). Active Portfolio Management (2a ed.). McGraw-Hill.",
 "Haugen, R. y Heins, A. (1975). Risk and the Rate of Return on Financial Assets. Journal of Financial and Quantitative Analysis, 10(5), 775-784.",
 "He, G. y Litterman, R. (1999). The Intuition Behind Black-Litterman Model Portfolios. Goldman Sachs Investment Management Research.",
 "Jegadeesh, N. y Titman, S. (1993). Returns to Buying Winners and Selling Losers. Journal of Finance, 48(1), 65-91.",
 "Markowitz, H. (1952). Portfolio Selection. Journal of Finance, 7(1), 77-91.",
 "McLean, R. y Pontiff, J. (2016). Does Academic Research Destroy Stock Return Predictability? Journal of Finance, 71(1), 5-32.",
 "Miffre, J. y Rallis, G. (2007). Momentum Strategies in Commodity Futures Markets. Journal of Banking & Finance, 31(6), 1863-1886.",
 "Moskowitz, T., Ooi, Y. y Pedersen, L. (2012). Time Series Momentum. Journal of Financial Economics, 104(2), 228-250."]:
    p = d.add_paragraph(ref); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6); p.runs[0].font.size = Pt(8.5)

# ================= ANEXOS =================
d.add_page_break(); H("Anexo A. Bitácora de auditoría del modelo anterior")
TB(["Gravedad", "Hallazgo", "Corrección"], [
    ["Alta", "Fórmulas sin valores guardados en el Excel", "Modelo v2 recalculado y guardado en Excel"],
    ["Alta", "Mezcla de dos arquitecturas (calce LDI de Educación vs cartera única) y notas 'PENDIENTE'", "Solo cartera única; textos depurados"],
    ["Alta", "Referencias a Betting Against Beta (BAB_EQ)", "Eliminadas; la estrategia es Low-Volatility"],
    ["Alta", "Meses jul-2018 y dic-2018 eliminados de las 6 series por un dato vacío de Low-Vol", "Ventana completa de 164 meses"],
    ["Alta", "Filtro de Sharpe de Momentum con volatilidad de 1 mes en el backtest y 12 meses en la operación", "Una sola regla: 12 meses"],
    ["Media", "Columna π desplazada una fila en Black-Litterman; μ_BL y pesos pegados como valores", "Fórmulas matriciales vivas (MMULT, MINVERSE)"],
    ["Media", "Benchmarks distintos entre hojas", "Un benchmark invertible por sleeve"],
    ["Media", "Universo de commodities ampliado mirando el backtest", "Declarado y controlado"],
    ["Baja", "Métricas de la misma estrategia con distintas convenciones", "Una convención: exceso sobre T-Bills"]], w=[1.8, 8.4, 5.8])
H("Anexo B. Reproducibilidad")
B(["codigo/config.py: todos los parámetros. codigo/estrategias.py: backtests. codigo/portafolio.py: Markowitz, Black-Litterman, benchmark y métricas.",
   "codigo/objetivos.py: flujos y TIR. codigo/montecarlo.py: simulación. codigo/validacion.py: robustez y walk-forward.",
   "codigo/run_all.py corre todo; codigo/excel_modelo.py genera el Excel; codigo/figuras.py e informe.py, este documento; codigo/senales.py, las señales de operación.",
   "La versión portada de los backtests reproduce exactamente los resultados del modelo anterior (diferencia máxima 10⁻¹⁶) antes de aplicar las correcciones."])
out = os.path.join(BASE, "informe", "Informe_Final_Nexum.docx"); d.save(out); print("OK", out)
